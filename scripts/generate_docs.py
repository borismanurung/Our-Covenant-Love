import json
from pathlib import Path
from docx import Document

# Load chapters summary
json_path = Path('C:/Users/SGU/Desktop/Boris/Ministry/Our Covenant Love/Web/src/data/chapters_summary.json')
chapters = json.loads(json_path.read_text(encoding='utf-8'))

output_dir = Path('C:/Users/SGU/Desktop/Boris/Ministry/Our Covenant Love/Web/Marriage For Life')
output_dir.mkdir(parents=True, exist_ok=True)

for ch in chapters:
    doc = Document()
    doc.add_heading(ch['title'], level=1)
    doc.add_paragraph(ch['summary'])
    filename = output_dir / f"Chapter {ch['id']} - {ch['title'].replace('/', '-')}.docx"
    doc.save(str(filename))
    print(f'Created {filename}')
